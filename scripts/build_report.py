"""Build the assignment report (reports/report.docx, and report.pdf when MS Word is available).

    python scripts/build_report.py            # docx + pdf (pdf via docx2pdf / MS Word)
    python scripts/build_report.py --no-pdf   # docx only

The script reads its numbers from the project outputs: reports/model_comparison.csv,
the MLflow DB, and screenshots/proof/05_ruff_pytest.txt. It also redraws the
architecture diagram. Screenshots are taken from screenshots/<name>.png when the file
exists; otherwise a clearly marked placeholder is inserted. After adding screenshots,
just re-run the script.
"""
from __future__ import annotations

import argparse
import re
import sqlite3
import sys
from datetime import date
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import pandas as pd  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "reports" / "figures"
SHOTS = ROOT / "screenshots"
OUT_DOCX = ROOT / "reports" / "report.docx"
OUT_PDF = ROOT / "reports" / "report.pdf"
REPO_URL = "<GitHub repository URL - to be added after the first push>"

MISSING_SHOTS: list[str] = []


# --------------------------------------------------------------------------- data
def comparison() -> pd.DataFrame:
    return pd.read_csv(ROOT / "reports" / "model_comparison.csv", encoding="utf-8")


def mlflow_runs() -> list[tuple[str, str, str]]:
    db = ROOT / "mlruns" / "mlflow.db"
    if not db.exists():
        return []
    con = sqlite3.connect(db)
    try:
        rows = con.execute(
            "SELECT r.name, r.status, COALESCE(t.value, '') FROM runs r "
            "LEFT JOIN tags t ON t.run_uuid = r.run_uuid AND t.key = 'selected' "
            "WHERE r.lifecycle_stage = 'active' ORDER BY r.start_time").fetchall()
    finally:
        con.close()
    return rows


def test_summary() -> tuple[str, str]:
    f = ROOT / "screenshots" / "proof" / "05_ruff_pytest.txt"
    if not f.exists():
        return "see pytest output", "ruff: see log"
    text = f.read_text(encoding="utf-8", errors="replace")
    m = re.search(r"(\d+) passed", text)
    ruff = "ruff: all checks passed" if "All checks passed" in text else "ruff: see log"
    return f"{m.group(1)} passed" if m else "see log", ruff


# --------------------------------------------------------------------------- diagram
def draw_architecture(path: Path) -> Path:
    fig, ax = plt.subplots(figsize=(13, 7.2))
    ax.set_xlim(0, 13)
    ax.set_ylim(0, 7.2)
    ax.axis("off")

    def box(x, y, w, h, title, body, color):
        ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.03,rounding_size=0.12",
                                    fc=color, ec="#34495e", lw=1.3))
        ax.text(x + w / 2, y + h - 0.22, title, ha="center", va="top", fontsize=10.5,
                weight="bold", color="#1b2631")
        ax.text(x + w / 2, y + h - 0.55, body, ha="center", va="top", fontsize=8.3,
                color="#1b2631", linespacing=1.35)

    def arrow(x1, y1, x2, y2, label="", style="-|>", ls="-", at=None):
        ax.add_patch(FancyArrowPatch((x1, y1), (x2, y2), arrowstyle=style, mutation_scale=14,
                                     lw=1.4, color="#2c3e50", linestyle=ls))
        if label:
            lx, ly = at or ((x1 + x2) / 2, (y1 + y2) / 2 + 0.14)
            ax.text(lx, ly, label, ha="center", fontsize=7.8, color="#2c3e50", style="italic",
                    bbox={"fc": "white", "ec": "none", "pad": 1.2, "alpha": 0.9})

    def zone(x, y, w, h, label, color):
        ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.02,rounding_size=0.2",
                                    fc=color, ec="#aab7b8", lw=1, ls="--", alpha=0.45))
        ax.text(x + 0.12, y + h - 0.1, label, fontsize=8.5, color="#566573", va="top",
                weight="bold")

    zone(0.15, 3.85, 8.1, 3.2, "BUILD & TRAIN  (dev laptop / GitHub Actions)", "#eaf2f8")
    zone(8.45, 3.85, 4.4, 3.2, "CI  (.github/workflows/ci.yml)", "#fef9e7")
    zone(0.15, 0.15, 12.7, 3.5, "SERVE & OBSERVE  (Docker Desktop: Kubernetes + compose)",
         "#e9f7ef")

    box(0.35, 5.0, 1.75, 1.6, "UCI repo",
        "processed.\ncleveland.data\n(fallback:\nucimlrepo id 45)", "#ffffff")
    box(2.4, 5.0, 1.75, 1.6, "Data prep", "download_data.py\nheart.data.clean\n-> heart_clean.csv",
        "#d6eaf8")
    box(4.45, 5.0, 1.85, 1.6, "Training", "heart.train\nPipeline + GridSearch\n5-fold strat. CV\n"
        "LogReg | RF | HGB", "#d6eaf8")
    box(6.6, 5.55, 1.5, 1.05, "MLflow", "mlruns/mlflow.db\n+ artifacts", "#fdebd0")
    box(6.6, 4.05, 1.5, 1.2, "Model", "model.joblib\nmlflow_model/", "#fdebd0")
    box(8.65, 4.15, 4.0, 2.45, "Pipeline on every push / PR",
        "checkout -> setup-python 3.12\n-> pip install -> ruff\n-> download data -> pytest\n"
        "-> python -m heart.train\n-> upload models/, reports/, mlruns/", "#fcf3cf")

    box(0.35, 0.45, 2.0, 2.3, "Docker image", "heart-api:0.2.0\npython:3.12-slim\nmulti-stage,"
        " non-root\nuvicorn :8000", "#d5f5e3")
    box(2.75, 0.45, 3.6, 2.3, "Kubernetes (docker-desktop)",
        "Deployment heart-api: 2 replicas\nreadiness/liveness -> /health\n"
        "100m/192Mi req, 500m/384Mi lim\nService LoadBalancer 8080 -> 8000", "#abebc6")
    box(6.75, 0.45, 1.9, 1.75, "Client", "localhost:8080\nPOST /predict\nGET /health\n"
        "load_test.py", "#ffffff")
    box(9.0, 1.65, 1.75, 1.3, "Prometheus", ":9090\nscrape /metrics\nevery 5 s", "#fadbd8")
    box(11.0, 1.65, 1.7, 1.3, "Grafana", ":3000\nprovisioned\ndashboard", "#fadbd8")
    box(9.0, 0.35, 3.7, 1.0, "Logs", "one JSON line per request\n(kubectl logs ... --prefix)",
        "#f2f3f4")

    arrow(2.1, 5.8, 2.4, 5.8)
    arrow(4.15, 5.8, 4.45, 5.8)
    arrow(6.3, 6.05, 6.6, 6.05, "")
    arrow(6.3, 5.2, 6.6, 4.65)
    arrow(7.35, 4.05, 1.35, 2.75, "docker build", ls="--", at=(3.2, 3.05))
    arrow(2.35, 1.6, 2.75, 1.6, "")
    arrow(6.75, 1.45, 6.35, 1.45, "HTTP", at=(6.55, 1.62))
    arrow(9.0, 2.5, 6.35, 2.5, "scrape /metrics via :8080", style="<|-", at=(7.7, 2.62))
    arrow(10.75, 2.3, 11.0, 2.3)
    arrow(6.35, 0.3, 9.0, 0.6, "stdout", ls=":", at=(8.85, 0.25))
    arrow(8.1, 5.6, 8.65, 5.6, "same\nsteps", style="<|-|>", ls=":")
    ax.set_title("Heart-disease risk service - end-to-end architecture", fontsize=13,
                 weight="bold", pad=6)
    fig.tight_layout()
    fig.savefig(path, dpi=160)
    plt.close(fig)
    return path


# --------------------------------------------------------------------------- docx helpers
def set_cell_bg(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


class Report:
    def __init__(self) -> None:
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.left_margin = sec.right_margin = Cm(2.2)
        sec.top_margin = sec.bottom_margin = Cm(2.0)
        normal = self.doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(10.5)
        normal.paragraph_format.space_after = Pt(5)
        self.fig_no = 0
        self.tab_no = 0

    def h1(self, text):
        self.doc.add_heading(text, level=1)

    def h2(self, text):
        self.doc.add_heading(text, level=2)

    def p(self, text, bold_lead: str | None = None):
        para = self.doc.add_paragraph()
        if bold_lead:
            para.add_run(bold_lead).bold = True
        para.add_run(text)
        return para

    def bullets(self, items):
        for it in items:
            if isinstance(it, tuple):
                para = self.doc.add_paragraph(style="List Bullet")
                para.add_run(it[0]).bold = True
                para.add_run(it[1])
            else:
                self.doc.add_paragraph(it, style="List Bullet")

    def code(self, text):
        para = self.doc.add_paragraph()
        run = para.add_run(text)
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        para.paragraph_format.left_indent = Cm(0.5)
        return para

    def caption(self, text):
        para = self.doc.add_paragraph()
        r = para.add_run(text)
        r.italic = True
        r.font.size = Pt(9)
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def figure(self, path: Path, caption: str, width_cm: float = 15.5):
        if not path.exists():
            self.p(f"[missing figure: {path.name}]")
            return
        self.doc.add_picture(str(path), width=Cm(width_cm))
        self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.fig_no += 1
        self.caption(f"Figure {self.fig_no}. {caption}")

    def screenshot(self, name: str, what: str, width_cm: float = 15.5):
        path = SHOTS / name
        if path.exists():
            self.figure(path, f"Screenshot - {what}", width_cm)
            return
        MISSING_SHOTS.append(name)
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.style = "Table Grid"
        cell = t.rows[0].cells[0]
        set_cell_bg(cell, "FFF4CE")
        para = cell.paragraphs[0]
        r = para.add_run(f"SCREENSHOT PLACEHOLDER - screenshots/{name}")
        r.bold = True
        r.font.color.rgb = RGBColor(0x9C, 0x57, 0x00)
        cell.add_paragraph(f"To capture: {what}")
        cell.add_paragraph("Save the image under that exact name and re-run "
                           "`python scripts/build_report.py`.")
        self.doc.add_paragraph()

    def table(self, df: pd.DataFrame, caption: str, font_pt: float = 8.5, widths=None):
        self.tab_no += 1
        self.caption(f"Table {self.tab_no}. {caption}")
        t = self.doc.add_table(rows=1, cols=len(df.columns))
        t.style = "Light Grid Accent 1"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, col in enumerate(df.columns):
            cell = t.rows[0].cells[i]
            cell.text = str(col)
            for r in cell.paragraphs[0].runs:
                r.bold = True
                r.font.size = Pt(font_pt)
        for _, row in df.iterrows():
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = str(v)
                for r in cells[i].paragraphs[0].runs:
                    r.font.size = Pt(font_pt)
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Cm(w)
        self.doc.add_paragraph()

    def page_break(self):
        self.doc.add_page_break()


# --------------------------------------------------------------------------- content
def build(rep: Report) -> None:
    cmp_df = comparison()
    runs = mlflow_runs()
    tests, ruff = test_summary()
    chosen = cmp_df.iloc[0]["model"]  # fallback if model_meta.json is absent
    meta_path = ROOT / "models" / "model_meta.json"
    if meta_path.exists():
        import json

        chosen = json.loads(meta_path.read_text(encoding="utf-8"))["model_name"]
    best = cmp_df.set_index("model").loc[chosen]

    d = rep.doc
    # ---- title
    t = d.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Heart-Disease Risk Classifier as a Monitored, Containerised API")
    r.bold = True
    r.font.size = Pt(20)
    st = d.add_paragraph()
    st.alignment = WD_ALIGN_PARAGRAPH.CENTER
    st.add_run("MLOps (AIMLCZG523) - Assignment 01\n").font.size = Pt(12)
    st.add_run("Jayanth Saish Pamarti  |  BITS Pilani WILP\n").font.size = Pt(11)
    st.add_run(f"Report generated {date.today():%d %B %Y}").font.size = Pt(9)
    rep.p(REPO_URL, bold_lead="Repository: ")
    rep.p("Local access: API http://localhost:8080 (Kubernetes LoadBalancer), "
          "Prometheus http://localhost:9090, Grafana http://localhost:3000.",
          bold_lead="Deployment: ")

    rep.h2("Summary")
    rep.p(
        "This project turns the UCI Cleveland heart-disease data into a small but complete "
        "MLOps system. A script downloads the data. An executed notebook explains it. One "
        "scikit-learn pipeline holds every preprocessing step, and three model families are "
        "tuned with stratified cross-validation and tracked in MLflow. The selected model is "
        "packaged with joblib and the MLflow format, tested with pytest, and wired into a "
        "GitHub Actions workflow. It is served by FastAPI in a Docker image, deployed with two "
        "replicas on Docker Desktop Kubernetes, and observed through structured logs, "
        f"Prometheus and Grafana. The final model is {chosen.replace('_', ' ')}, chosen on CV "
        f"recall: CV ROC-AUC {best['cv_roc_auc']}, test ROC-AUC {best['test_roc_auc']}.")
    rep.figure(FIG / "architecture.png", "System architecture: build/train, CI, serving and "
               "observability.", 16.5)

    # ---- 1 problem & data
    rep.h1("1. Problem and data")
    rep.p(
        "The goal is a binary classifier that estimates whether a patient has angiographic "
        "heart disease from 13 routinely collected attributes. The intended use is screening: "
        "the model flags patients who should get a closer look, so a missed positive (false "
        "negative) costs more than an unnecessary follow-up. That asymmetry drives the "
        "model-selection rule in Section 4.")
    rep.p(
        "The data is processed.cleveland.data from the UCI Machine Learning Repository: 303 "
        "patients and 14 columns, with no header and '?' marking missing values. "
        "scripts/download_data.py fetches it over HTTPS, falls back to the ucimlrepo package "
        "(dataset id 45) if the URL fails, writes data/raw/cleveland_raw.csv with column "
        "names, and then writes the cleaned data/processed/heart_clean.csv. The original "
        "target `num` grades disease from 0 to 4; it is binarised as num > 0 -> 1, which "
        "gives 139 positives (46%) against 164 negatives.")
    rep.table(pd.DataFrame([
        ["age, trestbps, chol, thalach, oldpeak", "continuous", "median impute + StandardScaler"],
        ["ca (vessels coloured, 0-3)", "count / ordinal", "median impute + StandardScaler"],
        ["cp, restecg, slope, thal", "nominal codes",
         "mode impute + OneHotEncoder(ignore unknown)"],
        ["sex, fbs, exang", "binary flags", "mode impute, passed through"],
        ["pct_max_hr, st_per_mmhg", "derived in pipeline", "median impute + StandardScaler"],
    ], columns=["columns", "type", "treatment in the pipeline"]), "Column roles.",
        widths=[6.0, 3.2, 6.8])

    # ---- 2 EDA
    rep.h1("2. Exploratory data analysis")
    rep.p("The full analysis is in notebooks/01_eda.ipynb, which is executed with nbconvert "
          "so its outputs are saved. All figures are written to reports/figures/.")
    rep.figure(FIG / "class_balance.png", "Original severity grades vs. the binary target. "
               "Grades 1-4 range from 13 to 55 patients each, too few to model separately.", 14)
    rep.p(
        "The classes are close to balanced, so I did not resample. Stratification at every "
        "split, plus class_weight as a tunable option, was enough.", bold_lead="Class balance. ")
    rep.figure(FIG / "numeric_hist_by_target.png", "Numeric features by outcome (each class "
               "normalised separately).", 15.5)
    rep.p(
        "Peak exercise heart rate (thalach) and ST depression (oldpeak) separate the classes "
        "the most. Diseased patients peak about 19 bpm lower (medians 142 vs 161), and their "
        "median ST depression is 1.4 mm against 0.2 mm. The number of coloured vessels (ca) is "
        "similarly informative. Resting blood pressure has identical medians (130) in both "
        "groups, and cholesterol differs by only about 15 mg/dl. In this already-referred "
        "cohort, the classic risk factors matter much less than the stress-test results.",
        bold_lead="Distributions. ")
    rep.figure(FIG / "categorical_target_rate.png", "Disease rate per category level "
               "(dashed line = overall 46%).", 16)
    rep.p(
        "Asymptomatic chest pain (cp = 4) carries a 73% disease rate, while typical or "
        "atypical angina sits at 18-30%. That looks counter-intuitive, but it reflects a "
        "referral population in which silent ischaemia was found through other tests. "
        "Exercise-induced angina (77% vs 31%) and a reversible thallium defect (76%) are "
        "strong signals. Fasting blood sugar is almost flat (45% vs 49%). restecg = 1 has only "
        "4 patients, so a CV fold can miss that level entirely, which is why the encoder "
        "ignores unknown categories.", bold_lead="Categorical levels. ")
    rep.figure(FIG / "correlation_heatmap.png", "Spearman correlation (lower triangle).", 12.5)
    rep.p(
        "The strongest links to the target are thal (0.52), ca (0.49), cp (0.47), exang (0.43), "
        "thalach (-0.42) and oldpeak (0.41). Among the features themselves, the largest pair "
        "is oldpeak-slope (0.59). No pair is collinear enough to justify dropping a feature, "
        "and L2 regularisation absorbs the rest.", bold_lead="Correlation. ")
    rep.figure(FIG / "age_sex_risk.png", "Disease rate by age band and sex (90% CI).", 11.5)
    rep.p(
        "In men, the rate rises steadily with age, from 33% to 75%. In women it stays at 6-12% "
        "until 56-65, where it jumps to 50%. Combining chest pain with the thallium result "
        "shows a similar interaction: asymptomatic patients with a reversible defect (78 "
        "patients) reach 91% (figure cp_thal_interaction.png in the repo). These interactions "
        "are why tree ensembles were included alongside the linear baseline.",
        bold_lead="Own analysis - interactions. ")

    # ---- 3 features & cleaning
    rep.h1("3. Cleaning and feature reasoning")
    rep.bullets([
        ("Missing values. ", "Only 6 rows are incomplete: 4 missing ca and 2 missing thal. "
         "They are kept and imputed inside the pipeline (median for ca, mode for thal) rather "
         "than dropped or filled in the CSV. Dropping them would lose 2% of a small dataset. "
         "Imputing inside the pipeline fits the imputer on training folds only, so CV and "
         "test scores have no leakage, and the deployed API can accept a request without ca "
         "or thal."),
        ("Domain validation. ", "Coded columns are checked against the UCI codebook, and any "
         "value outside it becomes NaN and goes through the same imputation path. No "
         "out-of-domain values occur in the file; the check protects future inputs."),
        ("Encoding. ", "Nominal codes (cp, restecg, slope, thal) are one-hot encoded. ca is "
         "kept numeric because the disease rate rises roughly monotonically with the number "
         "of vessels. Binary flags pass through unchanged."),
        ("Derived features. ", "pct_max_hr = thalach / (220 - age) expresses peak heart rate "
         "relative to the age-predicted maximum, separating 'low because older' from 'could "
         "not reach target rate'. st_per_mmhg = oldpeak / trestbps x 100 gives linear models "
         "an ST-depression x blood-pressure interaction. Both are computed by a "
         "FunctionTransformer inside the pipeline, so callers send only the 13 raw fields. "
         "The EDA checked them honestly: pct_max_hr alone scores AUC 0.70 against 0.75 for "
         "raw thalach, so it complements thalach rather than replacing it. Neither appears in "
         "the selected model's ten largest coefficients."),
    ])

    # ---- 4 models
    rep.h1("4. Model development and tuning")
    rep.p(
        "Each candidate is a Pipeline(prep -> clf), where prep runs the derived features "
        "followed by a ColumnTransformer (impute + scale | impute + one-hot | impute binary). "
        "The data gets a stratified 80/20 hold-out split with seed 523; the 61 test rows are "
        "never used for tuning or selection. Each model is tuned by GridSearchCV with "
        "StratifiedKFold(5, shuffle=True, seed 523), scored on accuracy, precision, recall, F1 "
        "and ROC-AUC, and refit on ROC-AUC. The grids are deliberately small, because with "
        "~240 training rows a wide grid mostly overfits the folds.")
    rep.table(pd.DataFrame([
        ["Logistic Regression (saga)", "C {0.01-3}, l1_ratio {0, 0.5, 1}, class_weight", "36"],
        ["Random Forest", "trees {200, 500}, max_depth {3, 5, None}, min_samples_leaf {1, 3, 5}, "
         "max_features {sqrt, 0.5}", "36"],
        ["HistGradientBoosting", "learning_rate {0.03, 0.1}, max_depth {2, 3}, "
         "max_iter {100, 300}, l2 {0, 1}", "16"],
    ], columns=["model", "grid", "configs x 5 folds"]), "Search spaces.", widths=[4.2, 10, 2.3])
    show = cmp_df[["model", "cv_accuracy", "cv_precision", "cv_recall", "cv_f1", "cv_roc_auc"]]
    rep.table(show, "Cross-validation results on the training split (mean +/- std over 5 "
              "folds, best grid point).", font_pt=8)
    show2 = cmp_df[["model", "test_accuracy", "test_precision", "test_recall", "test_f1",
                    "test_roc_auc", "best_params"]]
    rep.table(show2, "Held-out test results (n = 61, threshold 0.5) and selected "
              "hyper-parameters.", font_pt=7.5)
    rep.p(
        "Every model whose CV ROC-AUC is within 0.02 of the best is shortlisted, and the "
        "shortlisted model with the highest CV recall wins, because recall protects against "
        "missed patients. "
        f"{chosen.replace('_', ' ').title()} led on both CV ROC-AUC ({best['cv_roc_auc']}) and "
        f"CV recall ({best['cv_recall']}), so it was selected. It is also the easiest to "
        "explain to a clinician. Random Forest shows slightly higher test recall (0.857 vs "
        "0.821), but that is a difference of one patient out of 28 positives in a 61-row test "
        "set, and it was clearly weaker in CV. Choosing on the test set would have been "
        "selection bias.", bold_lead="Selection rule. ")
    rep.figure(FIG / "roc_test.png", "ROC curves of the three tuned models on the held-out "
               "test set.", 10.5)
    rep.figure(FIG / "confusion_logreg.png", "Confusion matrix of the selected model "
               "(test set).", 8)
    rep.figure(FIG / "feature_importance.png", "Largest standardised coefficients of the "
               "selected model. Red raises the risk.", 11)
    rep.p("The largest coefficients are the number of coloured vessels (+0.94), asymptomatic "
          "chest pain (+0.87), male sex (+0.75), a reversible thallium defect (+0.63) and "
          "exercise angina (+0.62); a normal thallium scan lowers risk (-0.58). This matches "
          "the EDA.")

    # ---- 5 tracking
    rep.h1("5. Experiment tracking (MLflow)")
    rep.p(
        "heart.train delegates logging to a small RunTracker interface. heart.tracking."
        "MlflowTracker implements it, and --no-mlflow swaps in a no-op tracker for unit "
        "tests. Each candidate becomes one run in the experiment heart-disease-classifier, "
        "which logs the best grid parameters, the CV mean/std and test metrics for all five "
        "scores, the ROC and confusion-matrix PNGs, and the fitted pipeline "
        "(mlflow.sklearn.log_model with an input example). The winning run is tagged "
        "selected=true together with the written rationale.")
    rep.p(
        "MLflow 3.17 rejects the plain ./mlruns file backend by default, so run metadata is "
        "stored in a local SQLite file (mlruns/mlflow.db) and artifacts in mlruns/artifacts/. "
        "Everything stays local, with no server. MLflow 3 also stores sklearn models in the "
        "skops format, which only loads allow-listed types; the tracker allows exactly the "
        "flagged types that come from the heart, sklearn and numpy packages.",
        bold_lead="Implementation notes. ")
    if runs:
        rep.table(pd.DataFrame(runs, columns=["run name", "status", "selected"]),
                  f"Runs currently in the tracking store ({len(runs)}).", widths=[5, 4, 3])
    rep.code("mlflow ui --backend-store-uri sqlite:///mlruns/mlflow.db   # http://127.0.0.1:5000")

    # ---- 6 packaging
    rep.h1("6. Packaging and reproducibility")
    rep.bullets([
        ("Model artefacts. ", "The complete fitted pipeline (imputation, derived features, "
         "encoding/scaling and the classifier) is saved as models/model.joblib, which the API "
         "loads, and as models/mlflow_model/ via mlflow.sklearn.save_model, together with "
         "MLmodel, conda.yaml, requirements and an input example. models/model_meta.json "
         "records the chosen model, its parameters, metrics and selection rationale."),
        ("Environment. ", "requirements.txt pins every direct dependency to the version "
         "verified in the project venv (Python 3.12.10; e.g. scikit-learn 1.9.1, pandas 3.0.6, "
         "mlflow 3.17.1). requirements-serve.txt is the smaller runtime set used by the Docker "
         "image."),
        ("Code layout. ", "An installable src/ package (pip install -e . via pyproject.toml). "
         "config.py is the single source of truth for paths, column roles and the seed; "
         "data, features, models, evaluate, train, tracking, predict, api and observability "
         "each do one job."),
        ("Determinism. ", "A fixed seed (523) is used for the split, the CV shuffling and the "
         "models. Re-running training reproduces the table in Section 4."),
    ])
    rep.code("py -3.12 -m venv .venv\n"
             ".\\.venv\\Scripts\\python.exe -m pip install -r requirements.txt\n"
             ".\\.venv\\Scripts\\python.exe -m pip install -e .\n"
             ".\\.venv\\Scripts\\python.exe scripts\\download_data.py\n"
             ".\\.venv\\Scripts\\jupyter.exe nbconvert --to notebook --execute --inplace "
             "notebooks\\01_eda.ipynb\n"
             ".\\.venv\\Scripts\\python.exe -m heart.train")

    # ---- 7 CI
    rep.h1("7. Automated testing and CI/CD")
    rep.p(f"The pytest suite currently reports {tests}, and {ruff}. It covers:")
    rep.bullets([
        ("Data: ", "target binarisation, '?' parsed as missing, rows kept, out-of-domain codes "
         "nulled, missing-column error, stratified split ratio."),
        ("Features: ", "derived-feature arithmetic, output shape and finiteness after "
         "imputation, unseen categories ignored, numeric columns standardised."),
        ("Model: ", "pipeline fit/predict_proba, the three model families and their grids, the "
         "selection rule (recall within the AUC band), and an end-to-end quick training run "
         "on the real data with an AUC floor of 0.75."),
        ("API: ", "predict function, /health (ok and degraded), /predict (valid input, missing "
         "ca/thal, 422 on invalid input), /metrics content, and the JSON request-log format."),
    ])
    rep.p(
        ".github/workflows/ci.yml runs on push, pull request and manual dispatch. Its steps "
        "are checkout, setup-python 3.12 with pip cache, install requirements and the package, "
        "ruff (GitHub annotations), download the data, pytest with JUnit XML, full training "
        "with the log captured, and upload of models/, reports/ and mlruns/ as a per-run "
        "artifact even if an earlier step failed. Any failing step fails the job.")
    rep.p(
        "The workflow is written and every step has been run locally, but it has not yet run "
        "on GitHub. git was not installed on the development laptop when this report was "
        "built, and the repository has not been pushed. The screenshot below will be added "
        "after the first push.", bold_lead="Status, stated honestly: ")
    rep.screenshot("05_ci_github_actions.png", "GitHub Actions run page for ci.yml with all "
                   "steps green and the uploaded artifact listed")

    # ---- 8 containerization
    rep.h1("8. Containerisation")
    rep.p(
        "The Dockerfile has two stages on python:3.12-slim. The builder creates /opt/venv, "
        "installs requirements-serve.txt and then the heart package (without dev tools). The "
        "runtime stage copies only that venv plus models/model.joblib, runs as uid 10001, "
        "sets HEART_MODEL_PATH, declares a HEALTHCHECK against /health, and starts uvicorn on "
        "port 8000. .dockerignore keeps .venv, mlruns, data, notebooks and tests out of the "
        "build context.")
    rep.p("The API contract: POST /predict accepts the 13 raw fields as JSON (ca and thal are "
          "optional; ranges are validated by pydantic, and invalid input returns 422) and "
          "returns prediction, label, probability_disease and confidence = max(p, 1 - p). "
          "GET /health reports model status, version and pod.")
    rep.code('POST /predict  samples/sample_request.json\n'
             '-> {"prediction":1,"label":"heart disease","probability_disease":0.9671,'
             '"confidence":0.9671}\n'
             'POST /predict  samples/sample_request_low_risk.json\n'
             '-> {"prediction":0,"label":"no heart disease","probability_disease":0.0243,'
             '"confidence":0.9757}')
    rep.screenshot("06_docker_build.png", "terminal: end of `docker build -t heart-api:0.2.0 .`"
                   " and `docker images heart-api`")
    rep.screenshot("06_docker_run_predict.png", "terminal: `docker run -d -p 8000:8000 "
                   "heart-api:0.2.0`, then curl /health and curl /predict with "
                   "samples/sample_request.json")

    # ---- 9 k8s
    rep.h1("9. Production deployment on Kubernetes")
    rep.p(
        "The target is the Docker Desktop Kubernetes cluster (context docker-desktop, node "
        "desktop-control-plane, v1.36.1). k8s/deployment.yaml runs 2 replicas of "
        "heart-api:0.2.0 with imagePullPolicy IfNotPresent, so the locally built image is "
        "used. It defines readiness (every 5 s) and liveness (every 15 s) probes on /health, "
        "requests of 100m CPU / 192Mi and limits of 500m / 384Mi, runs as non-root, and uses "
        "a RollingUpdate with maxUnavailable 0. The image went from 0.1.0 to 0.2.0 this way "
        "with no downtime. k8s/service.yaml exposes the pods through a LoadBalancer service.")
    rep.p(
        "The first deployment used service port 80. localhost:80 on the laptop is already "
        "published by an unrelated container (another local web application), so Docker "
        "Desktop's load-balancer proxy could not bind it. kubectl describe svc showed "
        "SyncLoadBalancerFailed, and http://localhost/ returned the other application's page. "
        "The service now uses port 8080 -> 8000, and http://localhost:8080 is served by both "
        "replicas (/predict requests appear in both pods' logs and metrics).",
        bold_lead="Port-80 conflict. ")
    rep.code("kubectl apply -f k8s/\nkubectl rollout status deploy/heart-api\n"
             "kubectl get pods,svc,deploy -o wide\ncurl http://localhost:8080/health")
    rep.screenshot("07_kubectl_get.png", "terminal: `kubectl get pods,svc,deploy -o wide` "
                   "showing 2/2 pods Running, the LoadBalancer on 8080 and image 0.2.0")
    rep.screenshot("07_k8s_predict.png", "browser: http://localhost:8080/docs with /predict "
                   "executed using the sample body and a 200 response")
    rep.screenshot("07_docker_desktop_k8s.png", "Docker Desktop Kubernetes view with the "
                   "cluster running and the heart-api pods")

    # ---- 10 monitoring
    rep.h1("10. Monitoring and logging")
    rep.bullets([
        ("Request logging. ", "An HTTP middleware writes one JSON line per request to stdout "
         "with ts, pod, request_id (also returned as the x-request-id header), method, path, "
         "status, latency_ms, client, prediction and probability. Kubernetes collects these, "
         "and kubectl logs -l app.kubernetes.io/name=heart-api --prefix shows them."),
        ("Metrics. ", "/metrics uses prometheus_client and exposes heart_api_requests_total "
         "{method, path, status}, the heart_api_request_duration_seconds histogram, "
         "heart_api_predictions_total{label}, a histogram of the predicted probability, and "
         "heart_api_model_loaded. Every series carries a pod label. Path labels use the route "
         "template, which keeps label cardinality bounded."),
        ("Stack. ", "monitoring/docker-compose.yml runs Prometheus 3.5 (scraping "
         "host.docker.internal:8080 every 5 s) and Grafana 12.1, with the datasource and the "
         "dashboard provisioned from files. The dashboard 'Heart API - model service health' "
         "has headline stats (model loaded, total predictions, share predicted positive "
         "compared with the 46% training base rate as a drift hint, error rate, p95 latency) "
         "and time series for request rate, p50/p95/p99 latency, error share by status, "
         "predictions per class, requests per replica, and the predicted-risk distribution."),
        ("Verification. ", "After about 120 requests from scripts/load_test.py (real patient "
         "rows plus ~6% deliberately invalid payloads), the Prometheus target was UP, queries "
         "returned per-class prediction counts, 200/422 status counts and a p95 /predict "
         "latency of about 25 ms. Grafana's API reported the datasource healthy and the "
         "dashboard present."),
        ("Limitation. ", "Prometheus reaches the pods through the LoadBalancer and reuses one "
         "keep-alive connection, so most scrapes land on the same replica. The pod label keeps "
         "the replicas' counters as separate series, but per-replica totals are incomplete. In "
         "a real cluster, Prometheus would run inside it with kubernetes_sd_configs; the pods "
         "already carry prometheus.io/* annotations for that."),
    ])
    rep.screenshot("08_api_logs.png", "terminal: `kubectl logs -l app.kubernetes.io/name="
                   "heart-api --prefix --tail=20` with JSON lines for /predict")
    rep.screenshot("08_metrics_endpoint.png", "browser: http://localhost:8080/metrics showing "
                   "the heart_api_* series")
    rep.screenshot("08_prometheus_targets.png", "browser: http://localhost:9090/targets with "
                   "the heart-api target UP")
    rep.screenshot("08_prometheus_query.png", "browser: Prometheus query "
                   "`sum by (label) (heart_api_predictions_total)` in graph view")
    rep.screenshot("08_grafana_dashboard.png", "browser: http://localhost:3000/d/heart-api-ops"
                   " after a fresh load test, last 15 minutes, all panels populated")

    # ---- 11 limitations
    rep.h1("11. Limitations and next steps")
    rep.bullets([
        "303 patients from one hospital (Cleveland, 1988) is a small, dated, referral-biased "
        "sample. The metrics will not transfer to general screening without external "
        "validation, for example on the Hungarian, Swiss and Long Beach UCI subsets.",
        "The 0.5 threshold is not tuned. For screening, a lower threshold chosen on CV to hit "
        "a target recall (for example 0.90) would better match the cost asymmetry, and the "
        "probabilities should be calibration-checked.",
        "CI has not yet executed on GitHub, and there is no CD step. Next steps are to push, "
        "add an image build-and-push job, and deploy by tag.",
        "Monitoring tracks service health and the prediction mix, not data drift on the "
        "inputs. A per-feature drift check against the training distribution (PSI or KS) "
        "would be the next metric to add.",
        "Run Prometheus inside the cluster with pod discovery to fix the single-connection "
        "scrape bias. Add an Ingress and HPA if the service ever needed real traffic.",
    ])

    rep.h1("12. Repository and deliverables")
    rep.p(REPO_URL, bold_lead="Repository link: ")
    rep.table(pd.DataFrame([
        ["scripts/download_data.py, data/", "data acquisition + cleaned dataset"],
        ["notebooks/01_eda.ipynb, reports/figures/", "EDA (executed) and all figures"],
        ["src/heart/", "pipeline, training, tracking, prediction, API, observability"],
        ["reports/model_comparison.{csv,md}", "model comparison"],
        ["models/", "model.joblib, mlflow_model/, model_meta.json"],
        ["tests/, .github/workflows/ci.yml", "pytest suite and CI workflow"],
        ["Dockerfile, requirements*.txt", "container image and pinned dependencies"],
        ["k8s/", "Deployment + LoadBalancer Service"],
        ["monitoring/", "Prometheus + Grafana (compose, config, dashboard JSON)"],
        ["screenshots/, screenshots/proof/", "screenshots and saved command outputs"],
    ], columns=["path", "content"]), "Repository map.", widths=[7, 9])


def to_pdf(docx: Path, pdf: Path) -> bool:
    try:
        from docx2pdf import convert
    except ImportError:
        print("docx2pdf not installed - skipping PDF")
        return False
    try:
        convert(str(docx), str(pdf))
        return pdf.exists()
    except Exception as exc:  # noqa: BLE001 - Word missing or COM failure
        print(f"PDF conversion failed ({exc!r}); docx only")
        return False


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--no-pdf", action="store_true")
    args = ap.parse_args()
    draw_architecture(FIG / "architecture.png")
    rep = Report()
    build(rep)
    rep.doc.save(OUT_DOCX)
    print(f"wrote {OUT_DOCX}")
    if not args.no_pdf and to_pdf(OUT_DOCX, OUT_PDF):
        print(f"wrote {OUT_PDF}")
    if MISSING_SHOTS:
        print("screenshot placeholders remaining:", ", ".join(MISSING_SHOTS))
    else:
        print("all screenshots embedded")
    return 0


if __name__ == "__main__":
    sys.exit(main())
